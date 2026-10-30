import csv
import io
import re
import zipfile

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db.models.functions import Lower

from openpyxl import load_workbook


User = get_user_model()

REQUIRED_COLUMNS = (
    "roll_number",
    "first_name",
    "last_name",
    "email",
)

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 12 * 1024 * 1024
MAX_ROWS = 500

ROLL_PATTERN = re.compile(
    r"^[A-Z0-9][A-Z0-9_-]{2,39}$"
)


def normalize_row(values, headings, row_number):
    errors = []

    if len(values) != len(headings):
        errors.append(
            "Column count does not match the header."
        )

    values = list(values)

    if len(values) < len(headings):
        values.extend(
            [""] * (len(headings) - len(values))
        )

    raw = dict(
        zip(
            headings,
            values[:len(headings)],
        )
    )

    roll_value = raw.get("roll_number", "")
    first_value = raw.get("first_name", "")
    last_value = raw.get("last_name", "")
    email_value = raw.get("email", "")

    if not isinstance(roll_value, str):
        errors.append(
            "Roll number must be stored as text."
        )

    roll = str(roll_value or "").strip().upper()
    first_name = str(first_value or "").strip()
    last_name = str(last_value or "").strip()
    email = str(email_value or "").strip().lower()

    if not ROLL_PATTERN.fullmatch(roll):
        errors.append(
            "Invalid roll number format."
        )

    if not first_name or len(first_name) > 150:
        errors.append(
            "First name is required (max 150 characters)."
        )

    if not last_name or len(last_name) > 150:
        errors.append(
            "Last name is required (max 150 characters)."
        )

    for name in (first_name, last_name):
        if name.startswith(("=", "+", "-", "@")):
            errors.append(
                "Formula-like name values are not allowed."
            )
            break

        if any(ord(character) < 32 for character in name):
            errors.append(
                "Control characters are not allowed in names."
            )
            break

    try:
        validate_email(email)
    except ValidationError:
        errors.append("Invalid email address.")

    if len(email) > 254:
        errors.append(
            "Email address is too long."
        )

    required_domain = getattr(
        settings,
        "STUDENT_EMAIL_DOMAIN",
        "",
    ).lower().strip()

    if required_domain and (
        email.rsplit("@", 1)[-1] != required_domain
    ):
        errors.append(
            "Email must belong to the configured university domain."
        )

    return {
        "row_number": row_number,
        "roll_number": roll,
        "first_name": first_name,
        "last_name": last_name,
        "email": email,
        "errors": errors,
    }


def parse_rows(iterator):
    try:
        header = next(iterator)
    except StopIteration:
        raise ValidationError(
            "The spreadsheet is empty."
        )

    headings = [
        str(value or "").strip().lower()
        for value in header
    ]

    if (
        len(headings) != len(REQUIRED_COLUMNS)
        or set(headings) != set(REQUIRED_COLUMNS)
    ):
        raise ValidationError(
            "Spreadsheet must contain exactly these "
            "columns: "
            + ", ".join(REQUIRED_COLUMNS)
        )

    rows = []

    for row_number, values in enumerate(
        iterator,
        start=2,
    ):
        if not any(
            value is not None and str(value).strip()
            for value in values
        ):
            continue

        if len(rows) >= MAX_ROWS:
            raise ValidationError(
                "Maximum 500 student rows per import."
            )

        rows.append(
            normalize_row(
                values,
                headings,
                row_number,
            )
        )

    if not rows:
        raise ValidationError(
            "No student records were found."
        )

    return rows


def parse_students_file(uploaded_file):
    payload = uploaded_file.read(
        MAX_FILE_BYTES + 1
    )

    if len(payload) > MAX_FILE_BYTES:
        raise ValidationError(
            "File exceeds the 2 MB limit."
        )

    filename = uploaded_file.name.lower()

    if filename.endswith(".csv"):
        try:
            decoded = payload.decode("utf-8-sig")

            reader = csv.reader(
                io.StringIO(decoded, newline="")
            )

            rows = parse_rows(reader)

        except (UnicodeDecodeError, csv.Error):
            raise ValidationError(
                "CSV must be a valid UTF-8 file."
            )

    elif filename.endswith(".xlsx"):
        try:
            # Reject unusually large compressed
            # workbooks before XML parsing.
            with zipfile.ZipFile(
                io.BytesIO(payload)
            ) as archive:

                infos = archive.infolist()

                if (
                    len(infos) > 200
                    or sum(
                        item.file_size
                        for item in infos
                    ) > MAX_UNCOMPRESSED_BYTES
                ):
                    raise ValidationError(
                        "Excel workbook is too large."
                    )

            workbook = load_workbook(
                io.BytesIO(payload),
                read_only=True,
                data_only=False,
                keep_links=False,
            )

            try:
                if len(workbook.worksheets) != 1:
                    raise ValidationError(
                        "Excel file must contain exactly one worksheet."
                    )

                worksheet = workbook.worksheets[0]

                if (
                    worksheet.max_column is not None
                    and worksheet.max_column > 8
                ):
                    raise ValidationError(
                        "Excel worksheet has too many columns."
                    )

                def spreadsheet_rows():
                    for row in worksheet.iter_rows():

                        if any(
                            cell.data_type == "f"
                            for cell in row
                        ):
                            raise ValidationError(
                                "Excel formulas are not allowed."
                            )

                        yield [
                            cell.value
                            for cell in row
                        ]

                rows = parse_rows(
                    spreadsheet_rows()
                )

            finally:
                workbook.close()

        except (
            zipfile.BadZipFile,
            OSError,
            ValueError,
            KeyError,
        ):
            raise ValidationError(
                "Invalid or damaged XLSX workbook."
            )

    else:
        raise ValidationError(
            "Unsupported file type."
        )

    return inspect_existing_accounts(rows)


def inspect_existing_accounts(rows):
    """
    Validate duplicates in the file and
    existing Django user records.

    Database checks will run again at confirmation.
    """

    roll_numbers = [
        row["roll_number"].lower()
        for row in rows
        if row["roll_number"]
    ]

    emails = [
        row["email"].lower()
        for row in rows
        if row["email"]
    ]

    existing_usernames = set(
        User.objects.annotate(
            normalized_username=Lower("username")
        )
        .filter(
            normalized_username__in=roll_numbers
        )
        .values_list(
            "normalized_username",
            flat=True,
        )
    )

    existing_emails = set(
        User.objects.exclude(email="")
        .annotate(
            normalized_email=Lower("email")
        )
        .filter(
            normalized_email__in=emails
        )
        .values_list(
            "normalized_email",
            flat=True,
        )
    )

    seen_rolls = set()
    seen_emails = set()

    for row in rows:
        roll = row["roll_number"].lower()
        email = row["email"].lower()

        if roll in seen_rolls:
            row["errors"].append(
                "Duplicate roll number in file."
            )

        if email in seen_emails:
            row["errors"].append(
                "Duplicate email in file."
            )

        if roll in existing_usernames:
            row["errors"].append(
                "An account with this roll number already exists."
            )

        if email in existing_emails:
            row["errors"].append(
                "This email is already associated with an account."
            )

        seen_rolls.add(roll)
        seen_emails.add(email)

    return rows